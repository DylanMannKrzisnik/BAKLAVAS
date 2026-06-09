"""
Aggregate METASPACE annotation CSVs into one table and flag lipid-related hits.

Usage:
    conda run -n nichecompass_liana python scripts/aggregate_metaspace_annotations.py \
        --sample SMA_v11l12-038-d1 \
        --input-dir /path/to/metaspace_output
"""

from __future__ import annotations

import argparse
import ast
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


LIPID_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

LIPID_PATTERNS = [
    r"sn-glycero-3-phospho",
    r"phosphatidyl",
    r"glycerophosph",
    r"lysophosphatid",
    r"cardiolipin",
    r"plasmalogen",
    r"\bceramide\b",
    r"hexosylceramide",
    r"glucosylceramide",
    r"galactosylceramide",
    r"lactosylceramide",
    r"sphingomyelin",
    r"sphingosine",
    r"sphingolipid",
    r"cerebroside",
    r"sulfatide",
    r"ganglioside",
    r"cholesterol",
    r"\bsterol\b",
    r"triglyceride",
    r"triacylglycerol",
    r"diacylglycerol",
    r"monoacylglycerol",
    r"\bfatty acid\b",
    r"\bcarnitine\b",
    r"acylcarnitine",
    r"prostaglandin",
    r"leukotriene",
    r"thromboxane",
    r"eicosanoid",
    r"phosphocholine\b",
    r"dioleoyl",
    r"dipalmitoyl",
    r"docosahexaenoyl",
    r"octadecenoyl",
    r"hexadecanoyl",
    r"tetracosanoyl",
    r"tricosanoyl",
    r"behenoyl",
    r"lignoceroyl",
    r"\bPC\(",
    r"\bPE\(",
    r"\bPI\(",
    r"\bPS\(",
    r"\bPG\(",
    r"\bPA\(",
    r"\bSM\(",
    r"\bCer\(",
    r"\bLPC\(",
    r"\bLPE\(",
    r"\bLPI\(",
    r"\bCPA\(",
    r"psychosine",
    r"dag\(",
    r"tag\(",
]

LIPID_DATABASE_PREFIXES = ("LipidMaps_", "SwissLipids_")

EXCLUDE_PATTERNS = [
    r"phosphoserine",
    r"phospho-?l-?aspart",
    r"phospho-?d-?aspart",
    r"phosphonato-?l-?aspart",
    r"phosphonatooxy",
    r"ribosyl",
    r"ribofuranosyl",
    r"phosphonooxybutyr",
    r"phosphonatooxybutyr",
    r"phosphono\)",
    r"5-?o-?phosphono",
    r"alpha-d-ribosyl",
]

LIPID_RES = [re.compile(p, re.IGNORECASE) for p in LIPID_PATTERNS]
EXCLUDE_RES = [re.compile(p, re.IGNORECASE) for p in EXCLUDE_PATTERNS]


def parse_list(val) -> list[str]:
    if pd.isna(val):
        return []
    if isinstance(val, list):
        return [str(x) for x in val]
    s = str(val).strip()
    if not s:
        return []
    try:
        out = ast.literal_eval(s)
        return [str(x) for x in out] if isinstance(out, list) else [str(out)]
    except (ValueError, SyntaxError):
        return [s]


def is_lipid_related(text: str) -> bool:
    if not text or not str(text).strip():
        return False
    t = str(text)
    if any(rx.search(t) for rx in EXCLUDE_RES):
        return False
    return any(rx.search(t) for rx in LIPID_RES)


def lipid_match_terms(text: str) -> str:
    if not text:
        return ""
    t = str(text)
    if any(rx.search(t) for rx in EXCLUDE_RES):
        return ""
    hits = [rx.pattern for rx in LIPID_RES if rx.search(t)]
    return "; ".join(hits)


def load_annotations(input_dir: Path, sample: str, fdr: float) -> pd.DataFrame:
    pattern = f"{sample}.*.fdr{fdr}.csv"
    paths = sorted(
        p
        for p in input_dir.glob(pattern)
        if "summary" not in p.name.lower() and ".aggregated." not in p.name
    )
    if not paths:
        raise FileNotFoundError(f"No files matching {pattern} in {input_dir}")

    frames = []
    for path in paths:
        db_label = path.name[len(sample) + 1 :].rsplit(".fdr", 1)[0]
        df = pd.read_csv(path)
        df["database"] = db_label
        df["source_csv"] = str(path)
        frames.append(df)

    out = pd.concat(frames, ignore_index=True)
    out["fdr_threshold"] = fdr
    return out


def annotate_lipids(df: pd.DataFrame) -> pd.DataFrame:
    names = df["moleculeNames"].map(parse_list)
    ids = df["moleculeIds"].map(parse_list)

    df = df.copy()
    df["molecule_name_primary"] = names.map(lambda xs: xs[0] if xs else "")
    df["molecule_names_all"] = names.map(lambda xs: " | ".join(xs))
    df["molecule_ids_all"] = ids.map(lambda xs: " | ".join(xs))
    df["n_molecule_names"] = names.map(len)

    lipid_by_name = names.map(lambda xs: any(is_lipid_related(x) for x in xs))
    lipid_by_id = df["molecule_ids_all"].map(
        lambda s: bool(
            re.search(r"LIPID|LM[A-Z0-9]+|SWISSLIPID|SLM:\d+", s, re.IGNORECASE)
        )
    )
    lipid_by_db = df["database"].map(
        lambda db: any(str(db).startswith(prefix) for prefix in LIPID_DATABASE_PREFIXES)
    )
    df["is_lipid_related"] = lipid_by_name | lipid_by_id | lipid_by_db
    df["lipid_match_terms"] = names.map(
        lambda xs: "; ".join(sorted({lipid_match_terms(x) for x in xs if lipid_match_terms(x)}))
    )
    return df


def build_display_table(df: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "database",
        "mz",
        "adduct",
        "ion",
        "formula",
        "molecule_name_primary",
        "molecule_names_all",
        "molecule_ids_all",
        "fdr",
        "msm",
        "moc",
        "intensity",
        "is_lipid_related",
        "lipid_match_terms",
        "rhoSpatial",
        "rhoSpectral",
        "offSample",
        "source_csv",
    ]
    present = [c for c in cols if c in df.columns]
    return (
        df[present]
        .sort_values(["is_lipid_related", "mz", "database"], ascending=[False, True, True])
        .reset_index(drop=True)
    )


def write_highlighted_xlsx(df: pd.DataFrame, path: Path) -> None:
    df.to_excel(path, index=False, sheet_name="annotations")
    wb = load_workbook(path)
    ws = wb["annotations"]
    header = [cell.value for cell in ws[1]]
    try:
        lipid_col = header.index("is_lipid_related") + 1
    except ValueError:
        lipid_col = None

    for row_idx in range(2, ws.max_row + 1):
        is_lipid = False
        if lipid_col is not None:
            val = ws.cell(row=row_idx, column=lipid_col).value
            is_lipid = val is True or str(val).lower() in {"true", "1"}
        if is_lipid:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).fill = LIPID_FILL

    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sample", default="SMA_v11l12-038-b1")
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=Path("/home/mcb/users/dmannk/BAKLAVA_base/outputs/metaspace_output"),
    )
    parser.add_argument("--fdr", type=float, default=0.2)
    parser.add_argument(
        "--output-prefix",
        type=Path,
        default=None,
        help="Defaults to <input-dir>/<sample>.aggregated.fdr<fdr>",
    )
    args = parser.parse_args()

    raw = load_annotations(args.input_dir, args.sample, args.fdr)
    annotated = annotate_lipids(raw)
    table = build_display_table(annotated)

    prefix = args.output_prefix or (
        args.input_dir / f"{args.sample}.aggregated.fdr{args.fdr}"
    )
    csv_path = Path(str(prefix) + ".csv")
    xlsx_path = Path(str(prefix) + ".xlsx")

    table.to_csv(csv_path, index=False)
    write_highlighted_xlsx(table, xlsx_path)

    n_total = len(table)
    n_lipid = int(table["is_lipid_related"].sum())
    n_unique_ions = table.drop_duplicates(subset=["mz", "adduct", "ion"]).shape[0]
    n_lipid_ions = (
        table[table["is_lipid_related"]]
        .drop_duplicates(subset=["mz", "adduct", "ion"])
        .shape[0]
    )

    print(f"Aggregated {n_total} annotations from {table['database'].nunique()} databases")
    print(f"  Unique ions (m/z + adduct + ion): {n_unique_ions}")
    print(f"  Lipid-related annotations: {n_lipid} ({n_lipid_ions} unique lipid ions)")
    print(f"Wrote: {csv_path}")
    print(f"Wrote: {xlsx_path} (lipid rows highlighted)")

    by_db = (
        table.groupby("database", as_index=False)
        .agg(
            n_annotations=("mz", "size"),
            n_lipid_related=("is_lipid_related", "sum"),
        )
        .sort_values("database")
    )
    print("\nPer database:")
    print(by_db.to_string(index=False))


if __name__ == "__main__":
    main()
